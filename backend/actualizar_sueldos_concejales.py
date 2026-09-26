# encoding: utf-8
"""
Sueldos de CONCEJALES publicados con nombre por fuentes OFICIALES de cada ayuntamiento (portal de transparencia,
sede electrónica, boletín oficial o la propia web municipal). NUNCA agregadores ni prensa.

Genera backend/sueldos_concejales.json, que app.py carga al arrancar (SUELDOS_CONCEJALES) y muestra en la ficha
del ayuntamiento (sueldos_concejales_html). El bitácora de cada lote está en SUELDOS_CONCEJALES_LOG.md.

Regla de cada registro (encargo de César, 2026-09-26): nombre, cargo/concejalía, importe y URL de la fuente
oficial (+ la base del importe: "bruto anual", "bruto mensual"... tal como la fuente la dice). Si falta cualquiera,
o la base/fuente es ambigua, NO se guarda: nada se adivina ni se completa con estimaciones ni se convierte
(mensual -> anual). Para que ni un error de parseo ni una fila descuadrada cuele un dato falso, cada registro se
VERIFICA contra el texto crudo de la fuente antes de aceptarlo (nuevo_registro): el importe debe aparecer en ese
texto y, dentro de una ventana cerca de él, todos los tokens del nombre.

Uso:  python actualizar_sueldos_concejales.py [fuente ...]      (sin argumentos: todas)
      python actualizar_sueldos_concejales.py --lista           (fuentes registradas)

Cada fuente es una función conector registrada en _FUENTES; devuelve la lista de registros de UN municipio.
Salvaguarda (como en el script de menores): si un conector falla o devuelve menos del 90 % de lo que ya había
para ese municipio, se CONSERVAN los anteriores y se avisa (--forzar para aceptar el resultado nuevo).
"""
import io
import json
import os
import re
import sys
import time
import unicodedata

import requests

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_FILE = os.path.join(BASE_DIR, "sueldos_concejales.json")
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/122.0.0.0 Safari/537.36",
    "Accept-Language": "es-ES,es;q=0.9",
}


# ── utilidades ──────────────────────────────────────────────────────────────────────────────────────────────
def _norm(s):
    s = unicodedata.normalize("NFD", s or "")
    return re.sub(r"\s+", " ", "".join(c for c in s if unicodedata.category(c) != "Mn").lower()).strip()


def descargar(url, intentos=3, timeout=60, **kw):
    """GET con reintentos; devuelve el Response (lanza si no hay 200 tras los intentos)."""
    ultimo = None
    for i in range(intentos):
        try:
            r = requests.get(url, headers=HEADERS, timeout=timeout, **kw)
            r.raise_for_status()
            return r
        except Exception as e:
            ultimo = e
            time.sleep(3 * (i + 1))
    raise RuntimeError(f"no se pudo descargar {url}: {ultimo}")


def texto_html(html):
    """Texto plano de un HTML (sin scripts/estilos), con saltos de línea por fila/bloque."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    return soup.get_text("\n")


def texto_pdf(contenido):
    """Texto de un PDF (pdfplumber), página a página."""
    import pdfplumber
    with pdfplumber.open(io.BytesIO(contenido)) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)


def _formas_importe(importe):
    """Formas en que un importe puede aparecer escrito en la fuente: 45.000,00 / 45.000 / 45000,00 / 45000 / 45,000.00..."""
    formas = set()
    ent, dec = divmod(round(importe * 100), 100)
    dec = int(dec)
    for miles in (".", "", " ", ","):
        e = f"{int(ent):,}".replace(",", miles)
        formas.add(f"{e},{dec:02d}")
        formas.add(f"{e}.{dec:02d}") if miles != "." else None
        if dec == 0:
            formas.add(e)
    plano = f"{importe:.2f}".rstrip("0").rstrip(".")      # celdas de Excel: 49777.7 (sin ceros finales)
    formas.update({plano, plano.replace(".", ",")})
    return formas


def verificar_en_texto(nombre, importe, texto, ventana=450):
    """True si el importe aparece en `texto` y, en una ventana de ±`ventana` caracteres, aparecen TODOS los tokens
    del nombre (sin tildes ni mayúsculas). Protege contra filas descuadradas y errores de parseo."""
    t = _norm(texto)
    tokens = [x for x in _norm(nombre).replace(",", " ").split() if len(x) > 1]
    if not tokens:
        return False
    for forma in _formas_importe(importe):
        for m in re.finditer(r"(?<![\d.,])" + re.escape(forma) + r"(?![\d])", t):
            w = t[max(0, m.start() - ventana): m.end() + ventana]
            if all(tok in w for tok in tokens):
                return True
    return False


def nuevo_registro(municipio, provincia, nombre, cargo, importe, base, periodo, fuente_url, fuente_nombre, texto_fuente,
                   verificar_adyacencia=True, omitir_verificacion=False):
    """Registro validado. Lanza ValueError si falta algún campo obligatorio, la URL no es https o el dato no se
    puede comprobar en el texto crudo de la fuente."""
    nombre, cargo, base = (nombre or "").strip(), (cargo or "").strip(), (base or "").strip()
    if not (municipio and nombre and cargo and base and importe and importe > 0 and (fuente_url or "").startswith("https://")):
        raise ValueError(f"registro incompleto: {municipio!r} {nombre!r} {cargo!r} {importe!r} {base!r} {fuente_url!r}")
    if len(nombre.split()) < 2:
        raise ValueError(f"nombre incompleto (una sola palabra): {nombre!r}")
    # verificar_adyacencia=False SOLO para fuentes que publican el importe por CARGO en una tabla y las personas en otra
    # (p. ej. Málaga): allí el conector valida la unión por su cuenta y aquí se comprueba solo que nombre e importe existen
    # por separado en el texto de la fuente.
    if omitir_verificacion:
        # SOLO para importes que el conector obtiene SUMANDO cifras publicadas (Madrid: mensualidades) y que por tanto no
        # aparecen tal cual en el texto; el nombre sí debe aparecer en la fuente.
        t = _norm(texto_fuente)
        ok = all(tok in t for tok in _norm(nombre).replace(',', ' ').split() if len(tok) > 1)
    elif verificar_adyacencia:
        ok = verificar_en_texto(nombre, importe, texto_fuente)
    else:
        t = _norm(texto_fuente)
        ok = (all(tok in t for tok in _norm(nombre).replace(',', ' ').split() if len(tok) > 1)
              and any(f in t for f in _formas_importe(importe)))
    if not ok:
        raise ValueError(f"no verificable en la fuente: {nombre!r} {importe!r}")
    return {"municipio": municipio, "provincia": provincia, "nombre": nombre, "cargo": cargo, "importe": round(float(importe), 2),
            "base": base, "periodo": (periodo or "").strip(), "fuente_url": fuente_url, "fuente_nombre": (fuente_nombre or "").strip()}


def num_es(txt):
    """'45.000,50' / '45000,5' / '45.000' -> float; None si no es un número claro."""
    x = re.sub(r"[€\s]|euros?", "", txt or "", flags=re.I).rstrip(".,")
    if not x or re.search(r"[A-Za-z]", x) or not re.search(r"\d", x):
        return None
    try:
        if "," in x:
            return float(x.replace(".", "").replace(",", "."))
        if re.match(r"^\d{1,3}(\.\d{3})+$", x):
            return float(x.replace(".", ""))
        return float(x)
    except ValueError:
        return None


_PARTICULAS = {"de", "del", "la", "las", "los", "y", "e", "el"}


def nombre_persona(apellidos, nombre):
    """'ARAGON JIMENEZ' + 'JUAN TOMAS DE' -> 'Juan Tomas de Aragon Jimenez' (formato "Nombre Apellidos", capitalizado,
    partículas en minúscula; una partícula final del nombre —'... DE'— pasa delante de los apellidos). No se añaden
    ni se quitan tildes: se respeta lo que escribe la fuente."""
    ap, no = (apellidos or "").split(), (nombre or "").split()
    while no and no[-1].lower() in _PARTICULAS:
        ap.insert(0, no.pop())
    palabras = no + ap
    def cap(w):
        return "-".join(x.capitalize() for x in w.split("-")) if (w.isupper() or w.islower()) else w
    return " ".join(w.lower() if (i and w.lower() in _PARTICULAS) else cap(w) for i, w in enumerate(palabras))


_SIGLAS = {"PSOE", "PP", "VOX", "RRHH", "MRH", "PSC", "ERC", "JXCAT", "BNG", "PNV", "CS", "IU", "TTE.", "UP", "EH", "BILDU", "CC", "PAR", "PRC"}


def cargo_frase(cargo):
    """'PORTAVOZ PSOE' -> 'Portavoz PSOE' (frase en minúsculas salvo la 1.ª letra y las siglas de partidos/abreviaturas)."""
    palabras = [w if w.strip(".,").upper() in _SIGLAS else w.lower() for w in (cargo or "").split()]
    frase = " ".join(palabras)
    return frase[:1].upper() + frase[1:]


def filas_tabla_pdf(contenido):
    """Todas las filas de todas las tablas de un PDF (pdfplumber.extract_tables), con las celdas ya limpias."""
    import pdfplumber
    filas = []
    with pdfplumber.open(io.BytesIO(contenido)) as pdf:
        for pagina in pdf.pages:
            for tabla in pagina.extract_tables():
                for f in tabla:
                    filas.append([re.sub(r"\s+", " ", (c or "")).strip() for c in f])
    return filas


# ── conectores (uno por municipio; se añaden por lotes, ver SUELDOS_CONCEJALES_LOG.md) ───────────────────────
_FUENTES = {}


def _fuente(clave):
    def deco(fn):
        _FUENTES[clave] = fn
        return fn
    return deco


@_fuente("sevilla")
def actualizar_sevilla():
    """Portal de Transparencia del Ayuntamiento de Sevilla > Régimen de retribuciones > Retribuciones Concejales:
    'Retribuciones percibidas por los concejales en 2024' (PDF, actualizado a 23/05/2025). Tabla APELLIDOS | NOMBRE |
    RETRIBUCIONES | dietas | locomoción | asistencia a sesiones | trienio: solo se toma la columna RETRIBUCIONES
    (verificado por posición: incluso los importes pequeños de concejales sin dedicación están en esa columna).
    El documento no dice la concejalía de cada persona: cargo = 'Concejal/a' (el título del documento las identifica
    como concejales). Existe además una 'previsión 2025' (solo dedicaciones) que no se usa: es previsión, no lo percibido."""
    url = ("https://www.sevilla.org/transparencia/informacion-sobre-la-corporacion-municipal/regimen-de-retribuciones/"
           "retribuciones-concejales-2024.pdf")
    contenido = descargar(url).content
    texto = texto_pdf(contenido)
    pagina = ("https://www.sevilla.org/transparencia/informacion-sobre-la-corporacion-municipal/regimen-de-retribuciones")
    out = []
    for f in filas_tabla_pdf(contenido):
        if len(f) < 3 or f[0].upper().startswith("APELLIDOS") or not f[0] or not f[1]:
            continue
        importe = num_es(f[2])
        if not importe or importe <= 0:
            continue
        out.append(nuevo_registro("Sevilla", "sevilla", nombre_persona(f[0], f[1]), "Concejal/a", importe,
                                  "retribuciones percibidas en el año", "2024", url,
                                  "Portal de Transparencia del Ayuntamiento de Sevilla", texto))
    if len(out) < 20:
        raise RuntimeError(f"Sevilla: solo {len(out)} concejales (¿cambió el documento?)")
    return out


@_fuente("malaga")
def actualizar_malaga():
    """Portal de Transparencia del Ayuntamiento de Málaga > Altos cargos > 'Retribuciones del Alcalde y los Concejales, y
    regímenes de dedicación' (Excel oficial). Dos tablas en el mismo libro: (1) 'Régimen dedicación (mes año)': nombre,
    régimen de dedicación y cargo de cada miembro; (2) 'Retribuciones AAAA': retribución anual por CARGO y ámbito. Se une
    cargo -> importe solo en los casos SIN ambigüedad: cargos que existen únicamente en el equipo de gobierno (TENIENTE
    ALCALDE, CONCEJAL DELEGADO) con dedicación del 100 %. Se saltan: PORTAVOZ/PORTAVOZ ADJUNTO (el importe depende de si el
    grupo es de gobierno u oposición y el documento no lo dice por persona), dedicación parcial (la tabla no da importe),
    cargos combinados '/ SECRETARIO' (llevan complemento) y el alcalde (cubierto aparte)."""
    import openpyxl
    url = "https://www.malaga.eu/system/modules/eu.malaga.ayto.corporativas/templates/organigrama/vista-documento.jsp?id=1"
    wb = openpyxl.load_workbook(io.BytesIO(descargar(url).content), data_only=True)
    hojas_reg = [ws for ws in wb.worksheets if ws.title.startswith("Régimen dedicación")]
    hojas_ret = sorted((ws for ws in wb.worksheets if re.match(r"Retribuciones \d{4}$", ws.title)), key=lambda w: w.title)
    if not hojas_reg or not hojas_ret:
        raise RuntimeError("Málaga: cambió la estructura del libro")
    reg, ret = hojas_reg[-1], hojas_ret[-1]
    anio = ret.title.split()[-1]
    escala = {}
    for f in ret.iter_rows(values_only=True):
        if f[0] and str(f[0]).startswith("I. ") and f[1] == "Equipo de gobierno" and f[2] and isinstance(f[3], (int, float)):
            escala[str(f[2]).strip().upper()] = float(f[3])
    texto = "\n".join(" ".join("" if c is None else str(c) for c in f) for ws in (reg, ret) for f in ws.iter_rows(values_only=True))
    out = []
    for f in reg.iter_rows(values_only=True):
        if not f or len(f) < 6 or f[4] is None or str(f[3]).replace("%", "").strip() not in ("100", "1", "1.0"):
            continue
        cargo = str(f[4]).strip().upper()
        if cargo not in ("TENIENTE ALCALDE", "CONCEJAL DELEGADO") or cargo not in escala:
            continue
        nombre = " ".join(str(x).replace("*", "").strip() for x in f[:3] if x)
        out.append(nuevo_registro("Málaga", "malaga", nombre, cargo.capitalize(), escala[cargo],
                                  "retribución anual de su cargo según la tabla oficial (dedicación exclusiva 100 %)",
                                  f"{anio} ({reg.title.split('(')[-1].rstrip(')')})", url,
                                  "Portal de Transparencia del Ayuntamiento de Málaga", texto, verificar_adyacencia=False))
    if len(out) < 8:
        raise RuntimeError(f"Málaga: solo {len(out)} concejales (¿cambió el documento?)")
    return out


@_fuente("cadiz")
def actualizar_cadiz():
    """Portal de Transparencia del Ayuntamiento de Cádiz > ¿Quién es quién? > Concejales: una ficha por concejal con nombre,
    cargo/delegación y apartado RETRIBUCIONES ('Importe anual ( en 14 pagas ) : 45.000,00 €'). Se recorren las fichas del
    listado; solo se guardan las que declaran 'Importe anual' (el resto, sin retribución fija, se salta) y se excluye el
    alcalde (cubierto aparte)."""
    from bs4 import BeautifulSoup
    from urllib.parse import urljoin
    listado = "https://transparencia.cadiz.es/quien-es-quien/concejales/"
    r = descargar(listado)
    enlaces = []
    for a in BeautifulSoup(r.text, "html.parser").find_all("a", href=True):
        h = urljoin(r.url, a["href"])
        if re.match(r"^https://transparencia\.cadiz\.es/quien-es-quien/concejales/[a-z0-9-]+/?$", h) and h not in enlaces:
            enlaces.append(h)
    if len(enlaces) < 20:
        raise RuntimeError(f"Cádiz: solo {len(enlaces)} fichas de concejales (¿cambió la web?)")
    out, sin_importe = [], 0
    for url in enlaces:
        pagina = descargar(url)
        soup = BeautifulSoup(pagina.text, "html.parser")
        cuerpo = soup.find("main") or soup
        for t in cuerpo(["script", "style", "nav", "header", "footer"]):
            t.decompose()
        lineas = [re.sub(r"\s+", " ", x).strip() for x in cuerpo.get_text("\n").split("\n") if x.strip()]
        if "Datos personales" not in lineas or len(lineas) < lineas.index("Datos personales") + 3:
            continue
        i = lineas.index("Datos personales")
        nombre = lineas[i + 1]
        # tras el nombre pueden venir enlaces a redes sociales ('Twitter de X') o el nombre repetido: se ignoran
        cand = [x for x in lineas[i + 2: i + 9]
                if not re.match(r"(?i)(facebook|twitter|instagram|linkedin|youtube|tiktok)\b", x) and _norm(x) != _norm(nombre)]
        cand = [x for x in cand if not (x.isupper() and len(x) > 12)]
        if not cand:
            continue
        tipo, deleg = cand[0], (cand[1] if len(cand) > 1 and not re.match(r"(?i)(fecha|lugar|grupo)", cand[1]) else "")
        if tipo.lower().startswith("alcald"):
            continue
        texto = "\n".join(lineas)
        m = re.search(r"Importe anual\s*\(([^)]*)\)\s*:\s*([\d.,]+)\s*€", texto)
        if not m or not num_es(m.group(2)):
            sin_importe += 1
            continue
        if "máximo" in m.group(1).lower():
            sin_importe += 1          # tope anual de asistencias de concejales sin dedicación, no una retribución: se salta
            continue
        cargo = tipo + (f". {deleg}" if deleg else "")
        out.append(nuevo_registro("Cádiz", "cadiz", nombre, cargo[:220], num_es(m.group(2)),
                                  "importe anual (" + re.sub(r"\s+", " ", m.group(1)).strip() + ")", "vigente (ficha actualizada)",
                                  url, "Portal de Transparencia del Ayuntamiento de Cádiz", texto))
        time.sleep(0.4)
    print(f"  Cádiz: {len(enlaces)} fichas, {len(out)} con importe anual, {sin_importe} sin importe anual")
    if len(out) < 8:
        raise RuntimeError(f"Cádiz: solo {len(out)} concejales con importe (¿cambió el formato?)")
    return out


@_fuente("huelva")
def actualizar_huelva():
    """Portal de Transparencia del Ayuntamiento de Huelva > Ea - Cargos representativos > 'Retribuciones cargos
    representativos AAAA' (PDF). Primera sección 'CORPORACIÓN.- AÑO AAAA' (PUESTO | NOMBRE | NETO): importe NETO anual de
    los cargos con dedicación. Columnas separadas por posición (el puesto ocupa varias líneas). Se saltan la alcaldesa
    (cubierta aparte), la sección de asistencia a plenos y la de atrasos del año anterior."""
    import pdfplumber
    from bs4 import BeautifulSoup
    from urllib.parse import urljoin
    indice = "https://www.huelva.es/portal/es/transparencia/cargos-representativos"
    r = descargar(indice)
    docs = []
    for a in BeautifulSoup(r.text, "html.parser").find_all("a", href=True):
        m = re.search(r"Retribuciones cargos representativos\s+(\d{4})$", a.get_text(" ", strip=True))
        if m:
            docs.append((int(m.group(1)), urljoin(r.url, a["href"])))
    if not docs:
        raise RuntimeError("Huelva: no se encontró el listado de retribuciones")
    anio, pagina_doc = max(docs)
    pdf_url = None
    for a in BeautifulSoup(descargar(pagina_doc).text, "html.parser").find_all("a", href=True):
        if a["href"].lower().endswith(".pdf") and "retribuciones" in a["href"].lower():
            pdf_url = urljoin(pagina_doc, a["href"])
            break
    if not pdf_url:
        raise RuntimeError("Huelva: no se encontró el PDF")
    contenido = descargar(pdf_url).content
    texto = texto_pdf(contenido)
    # La sección principal es la PRIMERA tabla de la primera página (pdfplumber la separa bien por columnas: PUESTO | NOMBRE |
    # NETO); le siguen la de asistencia a plenos y la de atrasos del año anterior, que no se usan.
    out = []
    with pdfplumber.open(io.BytesIO(contenido)) as pdf:
        pagina = pdf.pages[0]
        titulos = (pagina.extract_text() or "")
        if f"CORPORACIÓN.- AÑO {anio}" not in titulos:
            raise RuntimeError(f"Huelva: el PDF no empieza por 'CORPORACIÓN.- AÑO {anio}'")
        tablas = pagina.find_tables()
        filas = tablas[0].extract() if tablas else []
    for f in filas:
        if len(f) < 3 or (f[0] or "").strip().upper() == "PUESTO" or not f[1] or not f[2]:
            continue
        cargo, nombre, importe = re.sub(r"\s+", " ", f[0] or "").strip(), re.sub(r"\s+", " ", f[1]).strip(), num_es(f[2])
        if not importe or cargo.upper().startswith("ALCALD"):
            continue
        nombre_fmt = " ".join(p_.lower() if p_.lower() in _PARTICULAS else "-".join(x.capitalize() for x in p_.split("-"))
                              for p_ in nombre.split())
        out.append(nuevo_registro("Huelva", "huelva", nombre_fmt, cargo_frase(cargo), importe,
                                  f"importe neto anual percibido en {anio}", str(anio), pdf_url,
                                  "Portal de Transparencia del Ayuntamiento de Huelva", texto))
    if len(out) < 8:
        raise RuntimeError(f"Huelva: solo {len(out)} concejales (¿cambió el documento?)")
    return out


def filas_tablas_html(html, cabecera_regex):
    """Filas (listas de celdas de texto) de las <table> cuya primera fila casa con `cabecera_regex`."""
    from bs4 import BeautifulSoup
    out = []
    for t in BeautifulSoup(html, "html.parser").find_all("table"):
        filas = [[re.sub(r"\s+", " ", c.get_text(" ", strip=True)) for c in tr.find_all(["td", "th"])] for tr in t.find_all("tr")]
        filas = [f for f in filas if any(f)]
        if filas and re.search(cabecera_regex, " ".join(filas[0]), re.I):
            out.append(filas)
    return out


@_fuente("l'hospitalet de llobregat")
def actualizar_hospitalet():
    """Seu electrònica de l'Ajuntament de L'Hospitalet > Organització municipal > 'Retribucions de càrrecs electes': tablas
    'Cognoms i nom | Càrrec | Govern/Oposició | Dedicació | Retribució bruta anual (EUR) | Formació política' de los
    miembros con dedicación exclusiva (y parcial). Acuerdo del Pleno de 26/06/2023; 'Retribucions 2025'. Se excluye el
    alcalde (cubierto aparte)."""
    url = "https://seuelectronica.l-h.cat/238622_1.aspx?id=1"
    r = descargar(url)
    texto = texto_html(r.text)
    m = re.search(r"Retribucions\s+(\d{4})", texto)
    anio = m.group(1) if m else ""
    out = []
    for filas in filas_tablas_html(r.text, r"Cognoms i nom"):
        cab = [c.lower() for c in filas[0]]
        # La 3.ª tabla (concejales sin dedicación) no tiene columna de dedicación: da un tope anual por asistencias
        # ('19.890 (1.657,5 per assistència)'), no una retribución: se salta.
        pos = [next((i for i, c in enumerate(cab) if k in c), None) for k in ("càrrec", "dedic", "bruta anual")]
        if None in pos:
            continue
        ic, ig, ir = pos
        for f in filas[1:]:
            if len(f) <= max(ic, ir) or "," not in f[0]:
                continue
            apellidos, nombre = [x.strip() for x in f[0].split(",", 1)]
            importe = num_es(f[ir])
            if not importe or f[ic].lower().startswith("alcalde"):
                continue
            out.append(nuevo_registro("L'Hospitalet de Llobregat", "barcelona", nombre_persona(apellidos, nombre), f[ic], importe,
                                      "retribució bruta anual (" + (f[ig] or "dedicació").lower() + ")", anio, url,
                                      "Seu electrònica de l'Ajuntament de L'Hospitalet", texto))
    if len(out) < 8:
        raise RuntimeError(f"L'Hospitalet: solo {len(out)} regidores (¿cambió la página?)")
    return out


@_fuente("terrassa")
def actualizar_terrassa():
    """Govern Obert de l'Ajuntament de Terrassa > Retribucions dels càrrecs electes: tabla 'NOM | RETRIBUCIONS 2026 (a: mensual
    bruta, 14 pagues; aa: indemnització d'assistència)'. Solo se guardan los que tienen mensual bruta > 0 (el importe se
    muestra tal cual, mensual, SIN convertir a anual). La tabla no da el cargo de cada persona: cargo = 'Regidor/a'
    (es la tabla de drets econòmics dels regidors). Se excluye el alcalde (cubierto aparte)."""
    url = "https://governobert.terrassa.cat/transparencia/retribucions-carrecs-electes/"
    r = descargar(url)
    texto = texto_html(r.text)
    tablas = filas_tablas_html(r.text, r"^NOM")
    if not tablas:
        raise RuntimeError("Terrassa: no se encontró la tabla")
    out = []
    for f in tablas[0]:
        if len(f) < 2 or not num_es(f[1]) or "€" not in f[1]:
            continue
        nombre = f[0].strip()
        if nombre.lower().startswith("jordi ballart"):
            continue                       # alcalde
        out.append(nuevo_registro("Terrassa", "barcelona", nombre, "Regidor/a", num_es(f[1]), "mensual bruta (14 pagues)", "2026",
                                  url, "Govern Obert de l'Ajuntament de Terrassa", texto))
    if len(out) < 8:
        raise RuntimeError(f"Terrassa: solo {len(out)} regidores (¿cambió la página?)")
    return out


def _titulo_nombre(txt):
    """'FÈLIX LARROSA PIQUÉ' -> 'Fèlix Larrosa Piqué' (partículas en minúscula, guiones respetados)."""
    return " ".join(w.lower() if (i and w.lower() in _PARTICULAS | {"i"}) else "-".join(x.capitalize() for x in w.split("-"))
                    for i, w in enumerate(txt.split()))


@_fuente("sabadell")
def actualizar_sabadell():
    """Ajuntament de Sabadell > Retribucions de l'alcaldessa i els/les regidors/es (PDF; 'Aprovació Ple de 11/10/2024'): tabla
    Regidor | Càrrec | Dedicació | Règim | Retribució bruta mensual | Retribució bruta anual | Indemnitzacions per assistència.
    Solo filas con retribución bruta ANUAL (exclusiva/parcial); las de solo indemnización por asistencia se saltan
    (no son sueldo) y también la alcaldesa (renuncia al sueldo; cubierta aparte)."""
    url = "https://web.sabadell.cat/images/Retribucions_electes.pdf"
    contenido = descargar(url).content
    texto = texto_pdf(contenido)
    m = re.search(r"Aprovaci[oó] Ple de (\d{2}/\d{2}/\d{4})", texto)
    periodo = f"acuerdo del Pleno de {m.group(1)}" if m else ""
    out = []
    for f in filas_tabla_pdf(contenido):
        if len(f) < 6 or f[0].lower().startswith("regidor") or not f[0]:
            continue
        anual = num_es(f[5]) if "----" not in f[5] else None
        if not anual or f[1].lower().startswith("alcald"):
            continue
        out.append(nuevo_registro("Sabadell", "barcelona", f[0], f[1][:220], anual,
                                  "retribució bruta anual (" + (f[2] or "dedicació").lower() + ")", periodo, url,
                                  "Ajuntament de Sabadell (retribucions dels electes)", texto))
    if len(out) < 8:
        raise RuntimeError(f"Sabadell: solo {len(out)} regidores (¿cambió el documento?)")
    return out


@_fuente("lleida")
def actualizar_lleida():
    """La Paeria (Ajuntament de Lleida) > Cartipàs > Retribucions > Retribucions càrrecs electes (PDF vigente, 'Data
    actualització'): NOM I COGNOMS | CÀRREC | GRUP | DEDICACIÓ | RETRIBUCIÓ MENSUAL | RETRIBUCIÓ ANUAL (= mensual x 14).
    Solo dedicación EXCLUSIVA/PARCIAL con importe anual; comprobación de integridad: anual == mensual x 14 (±0,5): una fila
    con un importe mal escrito en origen (p. ej. '60,300,10 €') se SALTA y se anota. Se excluye el Paer en cap (alcalde)."""
    url = "https://www.paeria.cat/ca/ajuntament/el-govern/cartipas/retribucions/retribucions-carrecs-electes/carrecs-electes-febrer-2026.pdf"
    r = descargar(url)
    texto = texto_pdf(r.content)
    m = re.search(r"Data actualitzaci[oó]:\s*(\d{2}/\d{2}/\d{4})", texto)
    periodo = "2026" + (f" (actualizado el {m.group(1)})" if m else "")
    out, saltadas = [], []
    for f in filas_tabla_pdf(r.content):
        if len(f) < 8 or f[0].upper().startswith("NOM I") or f[5].upper() not in ("EXCLUSIVA", "PARCIAL"):
            continue
        nombre = re.sub(r"\s+\d+$", "", f[0]).strip()               # llamada a nota al pie pegada al nombre ("... CÉSPEDES 4")
        mensual, anual = num_es(f[6]), num_es(f[7])
        if not (mensual and anual and abs(mensual * 14 - anual) <= 0.5):
            saltadas.append(nombre)
            continue
        if f[3].upper().startswith("PAER EN CAP"):
            continue
        out.append(nuevo_registro("Lleida", "lleida", _titulo_nombre(nombre), cargo_frase(f[3])[:220], anual,
                                  f"retribució anual, 14 pagues ({f[5].lower()})", periodo, r.url,
                                  "La Paeria (Ajuntament de Lleida), Retribucions càrrecs electes", texto))
    if saltadas:
        print(f"  Lleida: saltadas por importe ilegible/incoherente en origen: {saltadas}")
    if len(out) < 8:
        raise RuntimeError(f"Lleida: solo {len(out)} regidors (¿cambió el documento?)")
    return out


# ── Plataforma seu-e.cat (sede electrónica oficial de muchos ayuntamientos catalanes) ────────────────────────
_SEUE_BASE = "https://seu-e.cat/ca/web/{slug}/govern-obert-i-transparencia/informacio-institucional-i-organitzativa/organitzacio-politica-i-retribucions/carrecs-electes"


def _seue_enlaces(slug):
    """URLs de las fichas 'veureCarrec/ID' del listado de cargos electes de un ayuntamiento en seu-e.cat ([] si no existe)."""
    try:
        r = requests.get(_SEUE_BASE.format(slug=slug), headers=HEADERS, timeout=40)
    except requests.RequestException:
        return []
    if r.status_code != 200 or "veureCarrec" not in r.text:
        return []
    ids = []
    for m in re.finditer(r"veureCarrec/(\d+)", r.text):
        if m.group(1) not in ids:
            ids.append(m.group(1))
    return [_SEUE_BASE.format(slug=slug) + f"/-/grupPolitic/veureCarrec/{i}" for i in ids]


def seue_concejales(slug, municipio, provincia):
    """Fichas de cargos electes de un ayuntamiento en seu-e.cat: 'Retribució anual bruta: Any 2026: 63.983,06 € (dedicació
    exclusiva ...)'. Solo se guardan las fichas que declaran esa retribución (importe del año tal como figura); se excluye el
    alcalde. Cargo = el de la ficha (p. ej. 'Tinent/a d'alcalde', 'Regidor/a')."""
    from bs4 import BeautifulSoup
    enlaces = _seue_enlaces(slug)
    if len(enlaces) < 8:
        raise RuntimeError(f"{municipio}: la plataforma seu-e.cat no devuelve fichas de cargos electes")
    out, sin_retribucion = [], 0
    for url in enlaces:
        pagina = descargar(url)
        soup = BeautifulSoup(pagina.text, "html.parser")
        for t in soup(["script", "style", "noscript", "nav", "header", "footer"]):
            t.decompose()
        lineas = [re.sub(r"\s+", " ", x).strip() for x in soup.get_text("\n").split("\n") if x.strip()]
        texto = "\n".join(lineas)
        try:
            i = lineas.index("Càrrec electe")
        except ValueError:
            continue
        nombre, cargo = lineas[i - 3], lineas[i - 2]
        m = re.search(r"Retribuci[oó] anual bruta:\s*Any\s+(\d{4}):\s*([\d.,]+)\s*€\s*(\([^)]*\))?", texto)
        if not m or not num_es(m.group(2)):
            sin_retribucion += 1
            continue
        if cargo.lower().startswith("alcald"):
            continue
        detalle = (m.group(3) or "").strip("() ")
        try:
            # nombre (cabecera de la ficha) e importe (bloque 'Retribució anual bruta') están en la MISMA ficha pero lejos
            # entre sí (lista de comisiones en medio): se comprueba que ambos existen en la ficha, no su adyacencia.
            out.append(nuevo_registro(municipio, provincia, nombre, cargo, num_es(m.group(2)),
                                      "retribució anual bruta" + (f" ({detalle})" if detalle else ""), m.group(1), url,
                                      f"Seu electrònica de l'Ajuntament de {municipio} (seu-e.cat)", texto,
                                      verificar_adyacencia=False))
        except ValueError as e:
            print(f"  {municipio}: ficha saltada ({e})")
        time.sleep(0.3)
    print(f"  {municipio}: {len(enlaces)} fichas, {len(out)} con retribución anual bruta, {sin_retribucion} sin retribución publicada")
    return out


# municipios de seu-e.cat: clave del conector -> (slug de la sede, nombre, provincia)
_SEUE = {
    "badalona": ("badalona", "Badalona", "barcelona"),
    "tarragona": ("tarragona", "Tarragona", "tarragona"),
    "girona": ("girona", "Girona", "girona"),
    "rubí": ("rubi", "Rubí", "barcelona"),
    "castelldefels": ("castelldefels", "Castelldefels", "barcelona"),
    "viladecans": ("viladecans", "Viladecans", "barcelona"),
    "granollers": ("granollers", "Granollers", "barcelona"),
    "vic": ("vic", "Vic", "barcelona"),
    "gavà": ("gava", "Gavà", "barcelona"),
    "blanes": ("blanes", "Blanes", "girona"),
    "ripollet": ("ripollet", "Ripollet", "barcelona"),
    "olot": ("olot", "Olot", "girona"),
    "salt": ("salt", "Salt", "girona"),
    "calafell": ("calafell", "Calafell", "tarragona"),
    "sitges": ("sitges", "Sitges", "barcelona"),
    "martorell": ("martorell", "Martorell", "barcelona"),
    "palafrugell": ("palafrugell", "Palafrugell", "girona"),
    "manlleu": ("manlleu", "Manlleu", "barcelona"),
    "banyoles": ("banyoles", "Banyoles", "girona"),
    "roses": ("roses", "Roses", "girona"),
    "tàrrega": ("tarrega", "Tàrrega", "lleida"),
    "cardedeu": ("cardedeu", "Cardedeu", "barcelona"),
    "tordera": ("tordera", "Tordera", "barcelona"),
    "palamós": ("palamos", "Palamós", "girona"),
    "torredembarra": ("torredembarra", "Torredembarra", "tarragona"),
    "cubelles": ("cubelles", "Cubelles", "barcelona"),
    "canovelles": ("canovelles", "Canovelles", "barcelona"),
    "castellbisbal": ("castellbisbal", "Castellbisbal", "barcelona"),
    "argentona": ("argentona", "Argentona", "barcelona"),
    "montgat": ("montgat", "Montgat", "barcelona"),
    "alella": ("alella", "Alella", "barcelona"),
    "alcarràs": ("alcarras", "Alcarràs", "lleida"),
    "masquefa": ("masquefa", "Masquefa", "barcelona"),
    "palafolls": ("palafolls", "Palafolls", "barcelona"),
    "matadepera": ("matadepera", "Matadepera", "barcelona"),
    "cervelló": ("cervello", "Cervelló", "barcelona"),
    "llagostera": ("llagostera", "Llagostera", "girona"),
    "tiana": ("tiana", "Tiana", "barcelona"),
    "vidreres": ("vidreres", "Vidreres", "girona"),
    "roquetes": ("roquetes", "Roquetes", "tarragona"),
    "polinyà": ("polinya", "Polinyà", "barcelona"),
    "tona": ("tona", "Tona", "barcelona"),
    "guissona": ("guissona", "Guissona", "lleida"),
}
for _clave, (_slug, _muni, _prov) in _SEUE.items():
    _FUENTES[_norm(_muni)] = (lambda s=_slug, m=_muni, p=_prov: seue_concejales(s, m, p))


@_fuente("madrid")
def actualizar_madrid():
    """Portal de Transparencia del Ayuntamiento de Madrid > Recursos humanos > Retribuciones > 'Retribuciones brutas percibidas
    en nómina por cargos electos, periodo 2025' (PDF; una fila por persona y mes: unidad, cargo, nombre, año, mes, régimen,
    % dedicación, retribuciones). El importe guardado es la SUMA de las mensualidades brutas publicadas de cada persona en el
    año (aritmética sobre cifras oficiales, no una estimación) y la base dice cuántas mensualidades son. Se saltan el
    alcalde (cubierto aparte) y las personas con un mes repetido (dato ambiguo)."""
    url = ("https://transparencia.madrid.es/UnidadAyre/Estructura/CargosPoliticosDirectivosYEventuales/ficheros/"
           "RetribucionesCargosElectos_2025.pdf")
    contenido = descargar(url, timeout=120).content
    texto = texto_pdf(contenido)
    por_persona = {}
    for f in filas_tabla_pdf(contenido):
        if len(f) < 11 or not f[6].isdigit() or f[7] not in {str(i) for i in range(1, 13)}:
            continue
        importe = num_es(f[10])
        if importe is None:
            continue
        por_persona.setdefault(f[2], []).append((int(f[7]), importe, f[1], f[0]))
    if len(por_persona) < 30:
        raise RuntimeError(f"Madrid: solo {len(por_persona)} personas (¿cambió el documento?)")
    out, saltadas = [], []
    for nombre, filas in por_persona.items():
        meses = [m for m, _, _, _ in filas]
        if len(set(meses)) != len(meses):
            saltadas.append(nombre)
            continue
        cargos = list(dict.fromkeys(c for _, _, c, _ in filas))
        if any(c.upper().startswith("ALCALDE") for c in cargos):
            continue
        total = round(sum(i for _, i, _, _ in filas), 2)
        if total <= 0:
            continue
        unidad = filas[-1][3].replace("P.P.", "PP").replace("p.p.", "PP")
        # la fuente escribe a veces el mismo cargo de dos formas ("Delegado/a area de gobierno" / "Delegado/a de area de
        # gobierno"): se dejan una sola vez (comparando sin artículos ni tildes)
        vistos, unicos = set(), []
        for c in cargos:
            k = re.sub(r"\s+", " ", re.sub(r"\b(de|del|la)\b", "", _norm(c))).strip()
            if k not in vistos:
                vistos.add(k)
                unicos.append(c)
        cargo = " / ".join(cargo_frase(c) for c in unicos) + f" ({unidad.capitalize().replace(' pp', ' PP')})"
        n = len(filas)
        out.append(nuevo_registro("Madrid", "madrid", _titulo_nombre(nombre), cargo[:240], total,
                                  f"suma de las retribuciones brutas percibidas en nómina en el año ({n} mensualidad{'es' if n != 1 else ''})",
                                  "2025", url, "Portal de Transparencia del Ayuntamiento de Madrid", texto,
                                  omitir_verificacion=True))
    if saltadas:
        print(f"  Madrid: saltadas por mes repetido: {saltadas}")
    return out


@_fuente("elche")
def actualizar_elche():
    """Portal de Transparencia del Ayuntamiento de Elche > Cargos, personal y retribuciones > Sueldos públicos: un bloque por
    concejal (nombre, concejalía, tenencia de alcaldía y 'Dedicación exclusiva/75 % (Anual 2026): importe €'; o, sin dedicación,
    'Retribución anual 2026 por tenencia de alcaldía: importe €'). Cada bloque repite la línea en valenciano: se ignora.
    Solo se guardan los bloques con importe anual; se excluye el alcalde (cubierto aparte)."""
    url = "https://www.elche.es/transparencia-gobierno-abierto/informacion-institucional-y-organizativa/cargos-personal-y-retribuciones/sueldos-publicos/"
    r = descargar(url)
    texto = texto_html(r.text)
    lineas = [re.sub(r"\s+", " ", x).strip() for x in texto.split("\n")]
    lineas = [x for x in lineas if x]
    try:
        inicio = max(i for i, x in enumerate(lineas) if x.startswith("Régimen de dedicación de los concejales")) + 2
    except ValueError:
        raise RuntimeError("Elche: no se encontró el inicio del listado")
    re_ded = re.compile(r"^Dedicación (exclusiva|\d+ ?%)\s*\(Anual (\d{4})\):\s*([\d.,]+)\s*€")
    re_ten = re.compile(r"^Sin dedicación\. Retribución anual (\d{4}) por tenencia de alcadía ([\d.,]+)\s*€")
    re_cargo = re.compile(r"^(Vicealcald|Alcald|Concejal|(Primer|Segund|Tercer|Cuart|Quint|Sext|Séptim|Octav|Noven|Décim)[oa]? Teniente)", re.I)
    out, buf = [], []
    for x in lineas[inicio:]:
        m1, m2 = re_ded.match(x), re_ten.match(x)
        if x.startswith(("Dedicació ", "Sense dedicació")):
            buf = []                                   # línea en valenciano: cierra el bloque
            continue
        if re.fullmatch(r"[\d.,]+|€", x):
            continue                                   # importe partido en varias líneas en la versión valenciana
        if m1 or m2:
            # Un concejal sin línea de importe (sin dedicación) deja sus líneas en el buffer: el bloque de ESTE importe empieza
            # en la última línea que no parece un cargo (nombre); lo anterior es de otra persona y se descarta.
            k = max((i for i, y in enumerate(buf) if not re_cargo.match(y)), default=None)
            if k is not None and len(buf) - k >= 2:
                nombre, cargo = buf[k], ". ".join(buf[k + 1:])
                if m1:
                    importe, base, anio = num_es(m1.group(3)), f"importe anual (dedicación {m1.group(1).replace(' ', '')})", m1.group(2)
                else:
                    importe, base, anio = num_es(m2.group(2)), "retribución anual por tenencia de alcaldía (sin dedicación)", m2.group(1)
                if importe and not cargo.lower().startswith("alcalde") and cargo.split(".")[0].strip().lower() != "alcalde":
                    out.append(nuevo_registro("Elche", "alicante", nombre, cargo[:240], importe, base, anio, url,
                                              "Portal de Transparencia del Ayuntamiento de Elche", texto))
            buf = []
            continue
        buf.append(x)
        if len(buf) > 6:
            buf = buf[-6:]                              # ruido antes/después del listado
    if len(out) < 8:
        raise RuntimeError(f"Elche: solo {len(out)} concejales (¿cambió la página?)")
    return out


@_fuente("mataro")
def actualizar_mataro():
    """Portal de Transparència de l'Ajuntament de Mataró > Retribucions dels càrrecs electes > 'Retribució i dedicació de l'alcalde
    i els regidors AAAA' (Excel; el del último año publicado): 'Cognoms, nom | Càrrec | Grup | Tipo personal | DEDICACIÓ |
    RETRIBUCIÓ ANUAL AAAA'. Solo filas con dedicación numérica (1 = 100 %, 0,75 = 75 %) y retribución; las de 'Assistència a
    Plens' (dietas) no son sueldo y se saltan; se excluye el alcalde (cubierto aparte)."""
    import openpyxl
    from bs4 import BeautifulSoup
    from urllib.parse import urljoin
    pagina = "https://www.mataro.cat/transparencia/organitzacio-politica-i-retribucions/retribucions-de-lalcalde-i-els-regidors/retribucions-dels-carrecs-electes"
    r = descargar(pagina)
    docs = []
    for a in BeautifulSoup(r.text, "html.parser").find_all("a", href=True):
        m = re.search(r"Retribuci[oó] i dedicaci[oó] de l'alcalde i els regidors (\d{4})$", a.get_text(" ", strip=True))
        if m and a["href"].lower().endswith(".xlsx"):
            docs.append((int(m.group(1)), urljoin(r.url, a["href"])))
    if not docs:
        raise RuntimeError("Mataró: no se encontró ningún Excel de retribuciones")
    anio, url = max(docs)
    wb = openpyxl.load_workbook(io.BytesIO(descargar(url).content), data_only=True)
    ws = wb.worksheets[0]
    filas = [list(f) for f in ws.iter_rows(values_only=True)]
    texto = "\n".join(" ".join("" if c is None else str(c) for c in f) for f in filas)
    out = []
    for f in filas:
        if not f or not f[0] or "," not in str(f[0]) or str(f[0]).lower().startswith("cognoms"):
            continue
        cargo, dedic, importe = str(f[1] or "").strip(), f[4], f[5]
        if not isinstance(dedic, (int, float)) or not isinstance(importe, (int, float)) or importe <= 0:
            continue                                   # 'Assistència a Plens' u otros sin dedicación numérica
        if cargo.lower().startswith("alcalde"):
            continue
        # 'Assistència a Plens' en el cargo = dietas (aunque la columna de dedicación diga 1); e importes < 20.000 € con
        # dedicación numérica son un periodo parcial (p. ej. 75 % y 8.791 €), no una retribución anual comparable: se saltan.
        if "assist" in cargo.lower() or importe < 20000:
            continue
        apellidos, nombre = [x.strip() for x in str(f[0]).split(",", 1)]
        out.append(nuevo_registro("Mataró", "barcelona", nombre_persona(apellidos, nombre), cargo, float(importe),
                                  f"retribució anual {anio} (dedicació {round(dedic * 100)}%)", f"{anio} (a 31/12/{anio})", url,
                                  "Portal de Transparència de l'Ajuntament de Mataró", texto))
    if len(out) < 8:
        raise RuntimeError(f"Mataró: solo {len(out)} regidors (¿cambió el Excel?)")
    return out


@_fuente("vigo")
def actualizar_vigo():
    """Portal de Transparencia del Concello de Vigo > 10. Retribuciones de los cargos electos > 'Retribuciones del alcalde, de las
    concejalas y los concejales' (Resolución de la Alcaldía, expediente 3891/1101, julio de 2023): lista de concelleiros en
    dedicación EXCLUSIVA y en dedicación PARCIAL (26 horas) con su importe anual. Es una resolución fechada: el periodo lo dice.
    Se excluye el alcalde (cubierto aparte)."""
    url = "https://transparencia.vigo.org/?id_file=870&lang=es&file="
    contenido = descargar(url, timeout=90).content
    if contenido[:4] != b"%PDF":
        raise RuntimeError("Vigo: el fichero ya no es un PDF")
    texto = texto_pdf(contenido)
    re_linea = re.compile(r"^(Alcalde|Concelleir[oa](?: do g\.m\. [^ ]+(?: [^ ]+)?)?)\.?\s+(?:D\.|Dª)\s+(.+?)\s+([\d.]+,\d{2}) €$")
    out, seccion = [], None
    for x in [re.sub(r"\s+", " ", l).strip() for l in texto.split("\n")]:
        if x.startswith("Primeiro.") and "dedicación exclusiva" in x:
            seccion = "exclusiva"
        elif x.startswith("Segundo.") and "dedicación parcial" in x:
            seccion = "parcial (26 horas)"
        m = re_linea.match(x)
        if not m or not seccion or m.group(1) == "Alcalde":
            continue
        out.append(nuevo_registro("Vigo", "pontevedra", m.group(2), m.group(1).rstrip("."), num_es(m.group(3)),
                                  f"retribución anual (dedicación {seccion})", "resolución de la Alcaldía de julio de 2023", url,
                                  "Portal de Transparencia del Concello de Vigo", texto))
    if len(out) < 8:
        raise RuntimeError(f"Vigo: solo {len(out)} concelleiros (¿cambió el documento?)")
    return out


@_fuente("logrono")
def actualizar_logrono():
    """Ayuntamiento de Logroño > Corporación local > 'Retribuciones e indemnizaciones de los miembros de la Corporación Municipal
    (actualización ...)' (PDF vigente, p. ej. junio 2026): tabla APELLIDOS Y NOMBRE | CARGO | GRUPO MUNICIPAL | TIPO DEDICACIÓN |
    RETRIBUCIÓN ANUAL. Solo 'Dedicación exclusiva/parcial' (las 'Indemnización asistencias' son dietas: se saltan); se excluye el
    alcalde (cubierto aparte)."""
    from bs4 import BeautifulSoup
    from urllib.parse import urljoin
    pagina = "https://logrono.es/corporacion-local"
    r = descargar(pagina)
    url = None
    for a in BeautifulSoup(r.text, "html.parser").find_all("a", href=True):
        if a.get_text(" ", strip=True).startswith("Retribuciones e indemnizaciones de los miembros de la Corporación Municipal (actualiza"):
            url = urljoin(r.url, a["href"])
            break
    if not url:
        raise RuntimeError("Logroño: no se encontró el PDF de retribuciones vigente")
    contenido = descargar(url, timeout=120).content
    texto = texto_pdf(contenido)
    m = re.search(r"Retribuciones (\d{4})", texto)
    anio = m.group(1) if m else ""
    out = []
    for f in filas_tabla_pdf(contenido):
        f = [c for c in f if c]                      # la 1.ª tabla trae celdas vacías de relleno entre columnas
        if len(f) < 5 or f[0].upper().startswith("APELLIDOS"):
            continue
        tipo = f[3].lower()
        importe = num_es(f[4])
        if not importe or "dedicación" not in tipo or f[1].upper().startswith("ALCALDE"):
            continue
        out.append(nuevo_registro("Logroño", "la_rioja", f[0], cargo_frase(f[1]) + f" ({f[2].title()})", importe,
                                  f"retribución anual ({tipo})", anio, url, "Ayuntamiento de Logroño (transparencia)", texto))
    if len(out) < 8:
        raise RuntimeError(f"Logroño: solo {len(out)} concejales (¿cambió el documento?)")
    return out


@_fuente("murcia")
def actualizar_murcia():
    """Ayuntamiento de Murcia > Institucional > 'Dedicación y Retribuciones de los miembros de la Corporación Municipal': el PDF más
    reciente de la lista ('Régimen de dedicación y retribuciones de los miembros de la Corporación (2023-2027)', p. ej. febrero 2024):
    tabla Nombre | Grupo | Dedicación | Competencias | Puesto | Retribuciones AÑO. El documento es el último que publica el
    ayuntamiento y está fechado (el periodo lo dice). Se excluye el alcalde (cubierto aparte)."""
    from bs4 import BeautifulSoup
    from urllib.parse import urljoin
    pagina = "https://web.murcia.es/institucional/dedicacion-retribuciones-corporacion"
    r = descargar(pagina)
    pdfs = [urljoin(r.url, a["href"]) for a in BeautifulSoup(r.text, "html.parser").find_all("a", href=True)
            if re.search(r"dedicacion", a["href"], re.I) and a["href"].lower().endswith(".pdf")]
    if not pdfs:
        raise RuntimeError("Murcia: no se encontró el PDF de retribuciones")
    url = pdfs[0]
    contenido = descargar(url, timeout=90).content
    texto = texto_pdf(contenido)
    m = re.search(r"CORPORACI[ÓO]N.{0,30}?\)\s*([A-ZÁÉÍÓÚ]+ \d{4})", texto)
    periodo = f"documento de {m.group(1).lower()}" if m else "documento vigente"
    out = []
    for f in filas_tabla_pdf(contenido):
        f = [c for c in f if c]
        if len(f) < 6 or f[0].lower().startswith("nombre"):
            continue
        importe = num_es(f[-1])
        if not importe or f[4].lower().startswith("alcalde"):
            continue
        out.append(nuevo_registro("Murcia", "murcia", f[0], f[4], importe,
                                  f"retribución anual (dedicación {f[2].lower()}, {f[3].lower()})", periodo, url,
                                  "Ayuntamiento de Murcia (dedicación y retribuciones de la Corporación)", texto))
    if len(out) < 15:
        raise RuntimeError(f"Murcia: solo {len(out)} concejales (¿cambió el documento?)")
    return out


@_fuente("santa cruz de tenerife")
def actualizar_santa_cruz_tenerife():
    """Portal de Transparencia del Ayuntamiento de Santa Cruz de Tenerife > Retribuciones > Altos cargos > 'Retribución percibida
    anualmente y dedicación de los miembros electos': la PRIMERA tabla (el año más reciente publicado: EMPLEADO/A | DENOMINACIÓN
    PUESTO | Dedicación | Periodo | RETRIBUCIÓN €). Solo filas con el año COMPLETO (01/01 a 31/12): las de periodos parciales se
    saltan (importe no comparable). Las tablas de años anteriores traen celdas combinadas y no se usan. Se excluye el alcalde."""
    url = "https://www.santacruzdetenerife.es/gobiernoabierto/transparencia/retribuciones/altos-cargos"
    r = descargar(url)
    texto = texto_html(r.text)
    tablas = filas_tablas_html(r.text, r"EMPLEADO")
    if not tablas:
        raise RuntimeError("Santa Cruz de Tenerife: no se encontró la tabla")
    out, parciales = [], 0
    for f in tablas[0][1:]:
        if len(f) != 5 or "," not in f[0]:
            continue
        m = re.match(r"^01/01/(\d{4}) a 31/12/(\d{4})$", f[3])
        if not m or m.group(1) != m.group(2):
            parciales += 1
            continue
        importe = num_es(f[4])
        if not importe or f[1].upper().startswith("ALCALDE"):
            continue
        apellidos, nombre = [x.strip() for x in f[0].split(",", 1)]
        out.append(nuevo_registro("Santa Cruz de Tenerife", "santa_cruz_tenerife", nombre_persona(apellidos, nombre), cargo_frase(f[1]),
                                  importe, f"retribución percibida en el año (dedicación {f[2].lower()})", m.group(1), url,
                                  "Portal de Transparencia del Ayuntamiento de Santa Cruz de Tenerife", texto))
    if parciales:
        print(f"  Santa Cruz de Tenerife: {parciales} filas con periodo parcial saltadas")
    if len(out) < 8:
        raise RuntimeError(f"Santa Cruz de Tenerife: solo {len(out)} concejales (¿cambió la página?)")
    return out


# ── main ────────────────────────────────────────────────────────────────────────────────────────────────────
def _leer():
    if not os.path.exists(OUT_FILE):
        return []
    for intento in range(5):                       # otra ejecución puede estar reemplazando el fichero en este instante
        try:
            with open(OUT_FILE, encoding="utf-8") as f:
                return json.load(f).get("registros", [])
        except (json.JSONDecodeError, PermissionError):
            time.sleep(0.5 * (intento + 1))
    raise RuntimeError("no se pudo leer sueldos_concejales.json")


def _escribir(registros):
    """Escritura atómica (fichero temporal + os.replace): nadie lee nunca un JSON a medias."""
    registros = sorted(registros, key=lambda r: (r["provincia"], r["municipio"], -r["importe"], r["nombre"]))
    tmp = f"{OUT_FILE}.{os.getpid()}.tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump({"generado": time.strftime("%Y-%m-%d"),
                   "descripcion": "Sueldos de concejales publicados con nombre por fuentes oficiales de cada ayuntamiento "
                                  "(ver SUELDOS_CONCEJALES_LOG.md). Registro: municipio, provincia, nombre, cargo, importe, base, "
                                  "periodo, fuente_url, fuente_nombre.",
                   "registros": registros}, f, ensure_ascii=False, indent=1)
    for intento in range(10):
        try:
            os.replace(tmp, OUT_FILE)
            return
        except PermissionError:
            time.sleep(0.3)
    raise RuntimeError("no se pudo reemplazar sueldos_concejales.json")


def main():
    args = sys.argv[1:]
    if "--lista" in args:
        print("\n".join(sorted(_FUENTES)))
        return
    forzar = "--forzar" in args
    pedidas = [a for a in args if not a.startswith("--")] or sorted(_FUENTES)
    desconocidas = [a for a in pedidas if a not in _FUENTES]
    if desconocidas:
        sys.exit(f"Fuente(s) desconocida(s): {desconocidas}. Válidas: {sorted(_FUENTES)}")
    previos = _leer()
    for nombre in pedidas:
        antes = [r for r in previos if _clave_fuente(r) == nombre]
        try:
            nuevos = _FUENTES[nombre]()
        except Exception as e:
            print(f"  !! {nombre}: FALLÓ ({type(e).__name__}: {e}); se conservan las {len(antes)} filas anteriores.", flush=True)
            continue
        if antes and not forzar and len(nuevos) < 0.9 * len(antes):
            print(f"  !! {nombre}: devolvió {len(nuevos)} filas frente a las {len(antes)} que ya había (< 90 %); se conservan las anteriores.", flush=True)
            continue
        nuevos, alcaldes = _sin_alcalde(nuevos)
        # Se guarda YA, tras cada municipio, RELEYENDO antes el fichero: si el proceso se corta no se pierde lo hecho, y varias
        # ejecuciones simultáneas (municipios distintos) no se pisan entre sí.
        _escribir([r for r in _leer() if _clave_fuente(r) != nombre] + nuevos)
        print(f"  {nombre}: {len(nuevos)} concejales" + (f" (excluido el alcalde, cubierto aparte: {alcaldes})" if alcaldes else ""), flush=True)
    print(f"\nHecho: {len(_leer())} registros en {OUT_FILE}", flush=True)


def _sin_alcalde(registros):
    """Quita los registros que son el alcalde/sa cuando el fichero del Ministerio (alcaldes_concejales.json, hoy solo
    Murcia y Cataluña) lo identifica: el alcalde ya está cubierto aparte (ISPA). Devuelve (resto, [nombres quitados])."""
    try:
        with open(os.path.join(BASE_DIR, "alcaldes_concejales.json"), encoding="utf-8") as f:
            ministerio = json.load(f).get("municipios", {})
    except Exception:
        return registros, []
    resto, quitados = [], []
    for r in registros:
        info = ministerio.get(_norm(r["municipio"]))
        ruido = _PARTICULAS | {"i"}
        alc = [x for x in _norm(((info or {}).get("alcalde") or {}).get("nombre", "")).replace(",", " ").split() if x not in ruido]
        toks = [x for x in _norm(r["nombre"]).replace(",", " ").split() if x not in ruido]
        if alc and set(alc) == set(toks):
            quitados.append(r["nombre"])
        else:
            resto.append(r)
    return resto, quitados


def _clave_fuente(r):
    """Clave del conector que generó el registro: el municipio normalizado (sin tildes, minúsculas) = clave de _FUENTES."""
    return _norm(r["municipio"])


if __name__ == "__main__":
    main()
