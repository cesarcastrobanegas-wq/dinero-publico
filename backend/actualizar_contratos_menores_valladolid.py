# encoding: utf-8
"""
Contratos MENORES del Ayuntamiento de Valladolid (~298.000 hab., Castilla y León).

Fuente OFICIAL: informes de "Contratación" en Excel del perfil de contratante (valladolid.gob.es/es/
perfil-contratante/contratos-menores-volumen-contratacion-tipo-procedimiento/ano-<año>). Cada año publica
varios ficheros AC ACUMULATIVOS (1er trimestre, luego 1er semestre, luego 1º+2º+3º trimestre, luego el
ejercicio completo) -- **solo hace falta el último de cada año** (el más completo), no todos.

**Cómo se encontraron los enlaces reales (2026-09-29)**: la página de cada año es estática y trae los `href`
directos para 2021-2024, pero 2025-2026 se sirven con una SPA en React que NO expone los enlaces de descarga en
el HTML inicial -- están en una SUB-página distinta y sí estática:
`.../ano-<año>/ayuntamiento-valladolid` (se encontró un `<a href>` real a esa subpágina dentro del propio HTML
de la página "SPA", en la lista de resultados -- ahí sí aparecen los `.ficheros/<id>-<nombre>.xlsx` reales).

**Estructura interna del Excel**: varias hojas; la relevante es "OPERACIONES SICALWIN" (el detalle fila a fila
de TODAS las operaciones de contratación del año, procedentes del programa de contabilidad municipal SICALWIN
-- no solo menores). Se filtra por la columna "PROCEDIMIENTO" == "Contratación menor" (los otros valores son
"Procedimiento abierto", "Acuerdo marco", "Procedimiento negociado sin publicidad" -- formales, no menores).
Columnas usadas: "Nº Operación" (clave única real), "Fecha" (fecha real, ya como `datetime` de Excel), "Nombre
Ter." (adjudicatario), "Texto Libre" (objeto), "Importe". **Sin NIF/CIF** -- esta fuente no lo publica en
absoluto.

Uso (desde backend/):
    python actualizar_contratos_menores_valladolid.py

Genera contratos_menores_valladolid.json.gz, que app.py carga al arrancar (ver
_cargar_contratos_menores_valladolid)."""
import gzip
import io
import json
import os
import re
import time
import urllib.request

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FICHERO = os.path.join(BASE_DIR, "contratos_menores_valladolid.json.gz")
MENORES_DESDE_FECHA = "2021-09-01"   # mismo corte que MENORES_DESDE_FECHA en app.py -- mantener sincronizado a mano

_BASE = ("https://www.valladolid.gob.es/es/perfil-contratante/"
         "contratos-menores-volumen-contratacion-tipo-procedimiento")
# año -> URL exacta del fichero ACUMULATIVO más completo disponible de ese año (recogidas a mano, 2026-09-29;
# para 2025/2026 la URL vive en la subpágina .../ano-<año>/ayuntamiento-valladolid, ver docstring). Cuando
# 2026 termine el año, esta URL habrá que actualizarla a mano por la del ejercicio completo.
FICHEROS = {
    2021: f"{_BASE}/ano-2021.ficheros/743925-CONTRATACION%20EJERCICIO%202021%20AYUNTAMIENTO%20VALLADOLID.xlsx",
    2022: f"{_BASE}/ano-2022.ficheros/852285-CONTRATACION%20EJERCICIO%202022%20AYUNTAMIENTO%20VALLADOLID.xlsx",
    2023: f"{_BASE}/ano-2023.ficheros/955845-CONTRATACION%20EJERCICIO%202023%20AYUNTAMIENTO%20VALLADOLID.xlsx",
    2024: f"{_BASE}/ano-2024.ficheros/1051832-CONTRATACION%20EJERCICIO%202024%20AYUNTAMIENTO%20VALLADOLID.xlsx",
    2025: (f"{_BASE}/ano-2025/ayuntamiento-valladolid.ficheros/"
           "1179354-CONTRATACION%20EJERCICIO%202025%20AYUNTAMIENTO%20VALLADOLID.xlsx"),
    2026: (f"{_BASE}/ano-2026/ayuntamiento-valladolid.ficheros/"
           "1207361-CONTRATACION%20PRIMER%20SEMESTRE%20EJERCICIO%202026%20AYUNTAMIENTO%20VALLADOLID.xlsx"),
}

_UA = "Mozilla/5.0 (compatible; dinero-publico-bot/1.0)"


def _get(url):
    req = urllib.request.Request(url, headers={"User-Agent": _UA})
    with urllib.request.urlopen(req, timeout=90) as r:
        return r.read()


def _limpiar(txt):
    return re.sub(r"\s+", " ", str(txt or "")).strip()


def _fecha_iso(valor):
    if hasattr(valor, "strftime"):
        return valor.strftime("%Y-%m-%d")
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$", str(valor or "").strip())
    if m:
        d, mo, y = m.groups()
        return f"{y}-{mo.zfill(2)}-{d.zfill(2)}"
    return ""


def _importe(valor):
    if isinstance(valor, (int, float)):
        return float(valor)
    limpio = str(valor or "").replace("€", "").strip().replace(".", "").replace(",", ".")
    try:
        return float(limpio)
    except ValueError:
        return 0.0


def _parsear(crudo, anio):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(crudo), data_only=True)
    if "OPERACIONES SICALWIN" not in wb.sheetnames:
        print(f"  !! {anio}: no se encontró la hoja 'OPERACIONES SICALWIN', fichero omitido "
              f"(hojas presentes: {wb.sheetnames})", flush=True)
        return []
    ws = wb["OPERACIONES SICALWIN"]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    cabecera = [_limpiar(c) for c in filas[0]]
    try:
        idx_op = cabecera.index("Nº Operación")   # "Nº Operación"
        idx_fecha = cabecera.index("Fecha")
        idx_proc = cabecera.index("PROCEDIMIENTO")
        idx_nombre = cabecera.index("Nombre Ter.")
        idx_texto = cabecera.index("Texto Libre")
        idx_importe = cabecera.index("Importe")
    except ValueError as e:
        print(f"  !! {anio}: cabecera de 'OPERACIONES SICALWIN' no reconocida ({e}) -- columnas vistas: "
              f"{cabecera}", flush=True)
        return []

    registros = []
    for fila in filas[1:]:
        if len(fila) <= max(idx_op, idx_fecha, idx_proc, idx_nombre, idx_texto, idx_importe):
            continue
        if _limpiar(fila[idx_proc]) != "Contratación menor":
            continue
        fecha = _fecha_iso(fila[idx_fecha])
        if not fecha or fecha < MENORES_DESDE_FECHA:
            continue
        registros.append({
            "id":               f"valladolid::{fila[idx_op]}",
            "municipio":         "Valladolid",
            "provincia":         "valladolid",
            "fuente":            "valladolid",
            "organisme":         "Ayuntamiento de Valladolid",
            "adjudicatari":      _limpiar(fila[idx_nombre]) or "No localizada",
            "nif":               "",   # no publicado por esta fuente (ver docstring)
            "import_num":        round(_importe(fila[idx_importe]), 2),
            "data_adjudicacio":  fecha,
            "tipus_contracte":   "",
            "descripcio":        _limpiar(fila[idx_texto]),
            "codi_cpv":          "",
            "exercici":          fecha[:4],
        })
    return registros


def _guardar(existentes):
    tmp = f"{FICHERO}.tmp"
    with gzip.open(tmp, "wt", encoding="utf-8") as f:
        json.dump({
            "generado": time.strftime("%Y-%m-%d"),
            "descripcion": ("Contratos menores del Ayuntamiento de Valladolid, filtrados de los informes "
                             "anuales de 'Contratación' (hoja OPERACIONES SICALWIN, PROCEDIMIENTO="
                             "'Contratación menor'). data_adjudicacio = 'Fecha' real. Sin NIF (no publicado "
                             "por esta fuente). Ver actualizar_contratos_menores_valladolid.py."),
            "registros": list(existentes.values()),
        }, f, ensure_ascii=False)
    os.replace(tmp, FICHERO)


def main():
    existentes = {}
    for anio, url in sorted(FICHEROS.items()):
        print(f"Descargando {anio}...", flush=True)
        try:
            crudo = _get(url)
        except Exception as e:
            print(f"  !! {anio}: error de descarga ({type(e).__name__}: {e})", flush=True)
            continue
        registros = _parsear(crudo, anio)
        for r in registros:
            existentes[r["id"]] = r
        print(f"  {anio}: {len(registros)} contratos en ventana", flush=True)
        _guardar(existentes)

    print(f"\nHecho: {len(existentes)} registros en {FICHERO}.")


if __name__ == "__main__":
    main()
